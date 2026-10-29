import json
import tempfile
import unittest
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch

from apps.control_api.documents import SourceRejected, acquire_run, get_coverage, search_evidence, render_pages


class DocumentTests(unittest.TestCase):
    def test_failed_pdf_child_is_scoped_to_page_and_next_page_continues(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.pdf"
            path.write_bytes(b"%PDF-synthetic")
            snapshot = {"files": [{"file_id": "f1", "path": str(path), "format": "pdf", "content_hash": sha256(path.read_bytes()).hexdigest()}]}
            outcomes = [{"ok": False, "error": "pdf_renderer_child_crashed"},
                        {"ok": True, "image_bytes": b"\x89PNG\r\n\x1a\nsynthetic"}]
            with patch("apps.control_api.pdf_renderer.render_pdf_page", side_effect=outcomes) as renderer:
                result = render_pages(snapshot, [{"file_id": "f1", "page": 1}, {"file_id": "f1", "page": 2}])
            self.assertEqual(renderer.call_count, 2)
            self.assertEqual(result["errors"], [{"file_id": "f1", "page": 1, "reason": "pdf_renderer_child_crashed"}])
            self.assertEqual([image["page"] for image in result["images"]], [2])

    def test_freezes_and_indexes_native_text_with_proof_required(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "daily.txt").write_text("2026-10-01 工作面 401 预警\n已安排撤人", encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"files": [{"path": "daily.txt", "file_id": "f1", "source_role": "actual_record"}], "source_complete": True, "scope_proof": True}), encoding="utf-8")
            snap = acquire_run({"run_id": "r1", "mine_id": "m1", "environment": "test", "rule_revision": "rev", "scope": {"manifest": str(root / "manifest.json")}}, {"input_root": str(root), "storage_dir": str(root / "store"), "allow_trusted_scope_proof": True})
            self.assertTrue(snap["source_complete"])
            self.assertEqual(search_evidence(snap, "预警")["total"], 1)
            self.assertEqual(get_coverage(snap)["readable"], 1)

    def test_production_manifest_flags_do_not_establish_scope_completeness(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "daily.txt").write_text("预警", encoding="utf-8")
            (root / "manifest.json").write_text(json.dumps({"files": [{"path": "daily.txt"}], "source_complete": True, "scope_proof": True}), encoding="utf-8")
            snap = acquire_run({"run_id": "r1", "mine_id": "m1", "environment": "production", "rule_revision": "rev", "scope": {"manifest": str(root / "manifest.json")}}, {"input_root": str(root), "storage_dir": str(root / "store"), "allow_trusted_scope_proof": True})
            self.assertFalse(snap["source_complete"])

    def test_docx_table_is_evidence_not_lost_body_text(self):
        from docx import Document
        from apps.control_api.documents import parse_file
        from hashlib import sha256
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.docx"
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "工作面"
            table.cell(0, 1).text = "记录"
            table.cell(1, 0).text = "401"
            table.cell(1, 1).text = "已安排撤人"
            doc.save(path)
            result = parse_file(path, file_id="f1", content_hash=sha256(path.read_bytes()).hexdigest(), source_role="actual_record")
            self.assertTrue(any("已安排撤人" in b["text"] and b["locator"].get("table") == 1 for b in result["blocks"]))

    def test_xlsx_formula_is_not_executed_or_published_as_value(self):
        from openpyxl import Workbook
        from apps.control_api.documents import parse_file
        from hashlib import sha256
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sample.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "=99+1"
            workbook.save(path)
            result = parse_file(path, file_id="f1", content_hash=sha256(path.read_bytes()).hexdigest(), source_role="actual_record")
            self.assertTrue(any("formula_without_cached_value" in c.get("reason", "") for c in result["coverage"]))
            self.assertFalse(any(b["text"] == "100" for b in result["blocks"]))

    def test_manifest_outside_root_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "root"
            root.mkdir()
            manifest = Path(tmp) / "manifest.json"
            manifest.write_text('{"files": []}', encoding="utf-8")
            with self.assertRaises(SourceRejected):
                acquire_run({"run_id": "r1", "mine_id": "m1", "rule_revision": "rev", "scope": {"manifest": str(manifest)}}, {"input_root": str(root)})


if __name__ == "__main__":
    unittest.main()

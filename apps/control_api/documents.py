"""Source freezing, native document extraction and evidence navigation.

The module intentionally stops before OCR. Pi asks for OCR only for selected
units through providers.py. Native evidence and model candidates remain
separate until the domain validator binds a candidate to source text.
"""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
from importlib.metadata import version
import math
import re
from collections import Counter
from pathlib import Path
from typing import Any


class SourceRejected(ValueError):
    pass


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _stable_evidence(content_hash: str, locator: dict[str, Any]) -> str:
    encoded = json.dumps([content_hash, locator], ensure_ascii=False, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _text_blocks(text: str, *, file_id: str, content_hash: str, source_role: str) -> list[dict[str, Any]]:
    blocks = []
    for index, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        locator = {"kind": "line", "line": index}
        blocks.append({"evidence_id": _stable_evidence(content_hash, locator), "document_id": file_id,
                       "document_version": content_hash, "content_hash": content_hash,
                       "locator": locator, "quote": line, "text": line, "source_role": source_role,
                       "unit_id": f"line:{index}"})
    return blocks


def parse_file(path: Path, *, file_id: str, content_hash: str, source_role: str) -> dict[str, Any]:
    if path.suffix.lower() in {".txt", ".md", ".csv", ".tsv", ".json"}:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
        return {"format": path.suffix[1:], "blocks": _text_blocks(text, file_id=file_id, content_hash=content_hash, source_role=source_role), "coverage": [{"unit_id": "document", "state": "readable" if text.strip() else "empty"}]}
    native = parse_document(path, file_id=file_id)
    blocks = []
    for block in native["blocks"]:
        locator = block["locator"]
        blocks.append({**block, "evidence_id": _stable_evidence(content_hash, locator), "document_id": file_id,
                       "document_version": content_hash, "content_hash": content_hash, "quote": block["text"],
                       "source_role": source_role, "unit_id": _unit_for_locator(locator)})
    return {"format": native["format"], "blocks": blocks, "coverage": native["coverage"], "issues": native["issues"]}


def _unit_for_locator(locator: dict[str, Any]) -> str:
    if "page" in locator:
        return f"page:{locator['page']}"
    if "sheet" in locator:
        return f"sheet:{locator['sheet']}:row:{locator.get('row', '')}"
    if "table" in locator:
        return f"table:{locator['table']}:row:{locator.get('row', '')}"
    if "paragraph" in locator:
        return f"paragraph:{locator['paragraph']}"
    return "document"


def _input_root(settings: dict[str, Any]) -> Path:
    raw = settings.get("input_root") or os.environ.get("COMPLIANCE_INPUT_ROOT")
    if not raw:
        raise SourceRejected("COMPLIANCE_INPUT_ROOT is required")
    root = Path(raw).resolve()
    if not root.is_dir():
        raise SourceRejected("configured input root does not exist")
    return root


def _parser_revision() -> str:
    return _sha(json.dumps({"source": LEGACY_PARSER_SHA256, "contract": "native-v3",
                            "dependencies": {name: version(name) for name in ("pypdf", "python-docx", "openpyxl")}}, sort_keys=True).encode())


def _parse_cached(path: Path, *, file_id: str, content_hash: str, source_role: str, storage: Path) -> tuple[dict, bool]:
    cache_key = _sha(json.dumps([content_hash, path.suffix.lower(), source_role, _parser_revision()]).encode())
    cache_path = storage / "native-cache" / f"{cache_key}.json"
    try:
        cached = json.loads(cache_path.read_text(encoding="utf-8"))
        parsed = cached["payload"]
        data_hash = _sha(json.dumps(parsed, ensure_ascii=False, sort_keys=True).encode())
        if cached["hash"] == data_hash and isinstance(parsed.get("blocks"), list) and isinstance(parsed.get("coverage"), list):
            for block in parsed["blocks"]:
                block["document_id"] = file_id
            return parsed, True
    except (OSError, ValueError, TypeError, KeyError, AttributeError):
        pass
    parsed = parse_file(path, file_id=file_id, content_hash=content_hash, source_role=source_role)
    if not any(unit.get("state") == "failed" for unit in parsed["coverage"]):
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        encoded = json.dumps(parsed, ensure_ascii=False, sort_keys=True)
        with tempfile.NamedTemporaryFile(dir=cache_path.parent, suffix=".part", mode="w", encoding="utf-8", delete=False) as stream:
            json.dump({"payload": parsed, "hash": _sha(encoded.encode())}, stream, ensure_ascii=False)
            temporary = Path(stream.name)
        os.replace(temporary, cache_path)
    return parsed, False


def _inventory(path: Path) -> dict:
    """Read container metadata only; text extraction belongs to a Pi-selected unit."""
    result = {"format": path.suffix.lower().lstrip(".") or "unknown", "blocks": [],
              "coverage": [{"unit_id": "document", "state": "not_indexed"}]}
    if result["format"] == "pdf":
        try:
            from pypdf import PdfReader
            with path.open("rb") as stream:
                reader = PdfReader(stream)
                if reader.is_encrypted:
                    result["inventory_issue"] = "encrypted_pdf"
                else:
                    result["page_count"] = len(reader.pages)
                    if result["page_count"]:
                        result["coverage"] = [{"unit_id": f"page:{page}", "state": "not_indexed"}
                                              for page in range(1, result["page_count"] + 1)]
        except Exception as exc:
            result["inventory_issue"] = f"pdf_inventory_failed:{type(exc).__name__}"
    elif result["format"] in {"png", "jpg", "jpeg", "bmp", "tiff", "tif", "webp"}:
        result["page_count"] = 1
    return result


def parse_frozen_file(snapshot: dict[str, Any], file_id: str, settings: dict[str, Any] | None = None) -> dict[str, Any]:
    """Parse one host-frozen object without mutating its immutable source inventory."""
    settings = settings or {}
    selected = [file for file in snapshot.get("files", []) if file.get("file_id") == file_id]
    if len(selected) != 1:
        raise SourceRejected("file outside frozen snapshot")
    file = selected[0]
    storage_value = settings.get("storage_dir") or snapshot.get("storage_root")
    if not storage_value:
        raise SourceRejected("frozen storage root is required")
    storage = Path(storage_value).resolve()
    path = Path(file.get("path", "")).resolve()
    if storage not in path.parents or not path.is_file():
        raise SourceRejected("frozen file outside configured storage")
    if path.stat().st_size > int(settings.get("max_file_bytes", 50 * 1024 * 1024)):
        raise SourceRejected("frozen file exceeds byte budget")
    if _sha(path.read_bytes()) != file.get("content_hash"):
        raise SourceRejected("frozen source checksum mismatch")
    aliases = [entry for entry in snapshot["files"] if entry.get("content_hash") == file["content_hash"]]
    origin = aliases[0]
    roles = {entry.get("source_role", "statement") for entry in aliases}
    role = file.get("source_role", "statement") if len(roles) == 1 else "statement"
    if role not in {"actual_record", "requirement", "statement", "regulation"}:
        raise SourceRejected("unknown source role")
    parsed, cached = _parse_cached(path, file_id=origin["file_id"], content_hash=file["content_hash"], source_role=role, storage=storage)
    owner = {key: snapshot[key] for key in ("run_id", "mine_id", "mine_code", "source_system", "environment", "rule_revision", "snapshot_id") if key in snapshot}
    owner.setdefault("mine_code", snapshot.get("mine_id"))
    owner.setdefault("source_system", "upstream")
    evidence = []
    for block in parsed["blocks"]:
        metadata = {**block.get("metadata", {}), **{key: origin[key] for key in ("report_date", "source_metadata") if key in origin}}
        evidence.append({**block, **owner, "file_id": origin["file_id"], "material_group": origin.get("material_group"),
                         "metadata": metadata, "source_file_ids": sorted(entry["file_id"] for entry in aliases)})
    coverage = [{"file_id": file_id, **unit} for unit in parsed["coverage"]]
    gaps = [f"{file_id}/{unit['unit_id']}:{unit.get('reason', unit['state'])}"
            for unit in coverage if unit.get("state") not in {"readable", "empty"}]
    if len(roles) > 1:
        gaps.append("same_content_source_role_conflict:" + file["content_hash"])
    return {"file_id": file_id, "evidence": evidence, "coverage": coverage, "coverage_gaps": gaps,
            "parser_revision": _parser_revision(), "cache_hit": cached}


def acquire_run(run: dict[str, Any], settings: dict[str, Any] | None = None) -> dict[str, Any]:
    settings = settings or {}
    scope = run.get("scope") or {}
    manifest_path = Path(scope.get("manifest") or run.get("manifest_path", ""))
    if not manifest_path.is_absolute():
        raise SourceRejected("manifest path must be absolute")
    root = _input_root(settings)
    manifest_path = manifest_path.resolve()
    if root not in manifest_path.parents:
        raise SourceRejected("manifest escapes configured input root")
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes.decode("utf-8-sig"))
    for field, expected in (("mine_id", run["mine_id"]), ("mine_code", run["mine_id"]), ("run_id", run["run_id"]), ("environment", run.get("environment", "test"))):
        if manifest.get(field) not in {None, expected}:
            raise SourceRejected("manifest ownership mismatch")
    entries = manifest.get("files")
    if not isinstance(entries, list):
        raise SourceRejected("manifest.files must be a list")
    storage = Path(settings.get("storage_dir", ".local/evidence"))
    storage.mkdir(parents=True, exist_ok=True)
    owner = {"run_id": run["run_id"], "mine_id": run["mine_id"], "environment": run.get("environment", "test"), "rule_revision": run["rule_revision"]}
    files, evidence, coverage, gaps = [], [], [], []
    gaps.extend("attachment_acquisition:" + str(failure.get("reason", "failed")) for failure in manifest.get("failures", []))
    seen_ids = set()
    for entry in entries:
        if not isinstance(entry, dict):
            raise SourceRejected("manifest file entry must be an object")
        rel = Path(str(entry.get("path", "")))
        source = (root / rel).resolve()
        if root not in source.parents or not source.is_file():
            gaps.append(f"missing_source:{entry.get('file_id') or rel}")
            continue
        max_bytes = int(settings.get("max_file_bytes", 50 * 1024 * 1024))
        if source.stat().st_size > max_bytes:
            gaps.append(f"file_too_large:{entry.get('file_id') or rel}")
            continue
        data = source.read_bytes()
        if len(data) > max_bytes:
            gaps.append(f"file_too_large:{entry.get('file_id') or rel}")
            continue
        content_hash = _sha(data)
        expected_hash = entry.get("content_hash")
        if expected_hash is not None and expected_hash != content_hash:
            gaps.append(f"source_hash_mismatch:{entry.get('file_id') or rel}")
            continue
        magic = {".pdf": b"%PDF-", ".docx": b"PK", ".xlsx": b"PK", ".xlsm": b"PK"}.get(source.suffix.lower())
        if magic and not data.startswith(magic):
            gaps.append(f"source_magic_mismatch:{entry.get('file_id') or rel}")
            continue
        file_id = str(entry.get("file_id") or content_hash[:24])
        if file_id in seen_ids:
            raise SourceRejected("duplicate file_id in manifest")
        seen_ids.add(file_id)
        role = str(entry.get("source_role") or entry.get("role") or "statement")
        if role not in {"actual_record", "requirement", "statement", "regulation"}:
            raise SourceRejected("unknown source role")
        frozen = storage / "objects" / f"{content_hash}{source.suffix.lower()}"
        frozen.parent.mkdir(parents=True, exist_ok=True)
        if frozen.exists():
            if _sha(frozen.read_bytes()) != content_hash:
                raise SourceRejected("frozen source checksum mismatch")
        else:
            with tempfile.NamedTemporaryFile(dir=frozen.parent, suffix=".part", delete=False) as stream:
                stream.write(data)
                temporary = Path(stream.name)
            os.replace(temporary, frozen)
        parsed = (_inventory(frozen) if settings.get("inventory_only") is True else
                  _parse_cached(frozen, file_id=file_id, content_hash=content_hash, source_role=role, storage=storage)[0])
        files.append({"file_id": file_id, "original_name": entry.get("original_name", source.name), "content_hash": content_hash,
                      "byte_size": len(data), "format": parsed["format"], "source_role": role, "material_group": entry.get("material_group"), "path": str(frozen),
                      **{key: entry[key] for key in ("report_date", "source_metadata", "original_url", "attachment_id", "attachment_url") if key in entry},
                      **{key: parsed[key] for key in ("page_count", "inventory_issue") if key in parsed}})
        for unit in parsed["coverage"]:
            coverage.append({"file_id": file_id, **unit})
            if unit.get("state") not in {"readable", "empty"}:
                gaps.append(f"{file_id}/{unit['unit_id']}:{unit.get('reason', unit['state'])}")
        evidence.extend({**owner, **block, "file_id": file_id, "material_group": entry.get("material_group")} for block in parsed["blocks"])
    unique_evidence = {}
    for block in evidence:
        prior = unique_evidence.get(block["evidence_id"])
        if prior is None:
            block["source_file_ids"] = [block["file_id"]]
            unique_evidence[block["evidence_id"]] = block
        else:
            prior["source_file_ids"] = sorted(set(prior.get("source_file_ids", [prior["file_id"]]) + [block["file_id"]]))
            if prior["source_role"] != block["source_role"]:
                prior["source_role"] = "statement"
                gaps.append("same_content_source_role_conflict:" + block["evidence_id"])
    evidence = list(unique_evidence.values())
    manifest_hash = _sha(manifest_bytes)
    explicit_complete = (settings.get("allow_trusted_scope_proof") is True and run.get("environment") != "production"
                         and manifest.get("source_complete") is True and manifest.get("scope_proof") is True)
    source_complete = bool(explicit_complete and files and not gaps and not settings.get("inventory_only"))
    snapshot_id = _sha(json.dumps({**owner, "manifest_hash": manifest_hash, "files": files}, ensure_ascii=False, sort_keys=True).encode())
    return {**owner, "snapshot_id": snapshot_id, "manifest_hash": manifest_hash, "profile": "native-input", "files": files,
            "evidence": evidence, "coverage": coverage, "coverage_gaps": list(dict.fromkeys(gaps)),
            "source_complete": source_complete, "observation": {"kind": "frozen-manifest", "scope_proof": explicit_complete},
            "storage_root": str(storage.resolve()),
            "scope_proof": run.get("scope_proof_receipt") or (manifest.get("scope_proof") if isinstance(manifest.get("scope_proof"), dict) else None),
            "parser_revision": LEGACY_PARSER_SHA256}


def search_evidence(snapshot: dict[str, Any], query: str = "", *, file_id: str | None = None, offset: int = 0, limit: int = 20) -> dict[str, Any]:
    if type(offset) is not int or offset < 0 or type(limit) is not int or not 1 <= limit <= 100:
        raise SourceRejected("invalid pagination")
    if len(query) > 2000:
        raise SourceRejected("search query too long")
    def terms(text):
        output = []
        for token in re.findall(r"[a-z0-9_.-]+|[\u3400-\u9fff]+", text.casefold()):
            output.extend([token] if len(token) < 2 or not re.match(r"[\u3400-\u9fff]", token) else [token[i:i + 2] for i in range(len(token) - 1)])
        return output
    rows = [row for row in snapshot.get("evidence", []) if not file_id or row.get("file_id") == file_id or file_id in row.get("source_file_ids", [])]
    query_terms = set(terms(query))
    if query.strip():
        bags = [Counter(terms(str(row.get("text", "")))) for row in rows]
        mean_length = sum(map(lambda bag: sum(bag.values()), bags)) / max(len(bags), 1) or 1
        counts = {term: sum(term in bag for bag in bags) for term in query_terms}
        ranked = []
        for row, bag in zip(rows, bags):
            length = sum(bag.values())
            score = sum(math.log(1 + (len(bags) - counts[term] + .5) / (counts[term] + .5))
                        * bag[term] * 2.2 / (bag[term] + 1.2 * (.25 + .75 * length / mean_length))
                        for term in query_terms if bag[term])
            if query.casefold().strip() in str(row.get("text", "")).casefold():
                score += 2
            if score > 0:
                ranked.append({**row, "score": round(score, 6)})
        rows = sorted(ranked, key=lambda row: (-row["score"], row["evidence_id"]))
    return {"items": rows[offset:offset + limit], "total": len(rows), "offset": offset, "limit": limit,
            "has_more": offset + limit < len(rows), "search_mode": "local-bm25-cjk-bigrams", "scope": "frozen_run_only",
            "absence_proven": False}


def get_coverage(snapshot: dict[str, Any]) -> dict[str, Any]:
    coverage = snapshot.get("coverage", [])
    return {"total": len(coverage), "readable": sum(item.get("state") in {"readable", "empty"} for item in coverage), "items": coverage, "gaps": snapshot.get("coverage_gaps", [])}


def render_pages(snapshot: dict[str, Any], pages: list[dict[str, Any]], *, max_pages: int = 16, dpi: int = 120) -> dict[str, Any]:
    """Render only Pi-selected PDF pages; renderer is an optional app dependency."""
    if len(pages) > max_pages:
        raise SourceRejected("render page budget exceeded")
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:
        return {"images": [], "errors": [{"reason": "page_renderer_unavailable", "detail": str(exc)}]}
    files = {item["file_id"]: item for item in snapshot.get("files", [])}
    images, errors = [], []
    for selected in pages:
        file = files.get(selected.get("file_id"))
        page_number = int(selected.get("page", 0))
        image_format = str((file or {}).get("format", "")).lower() in {"png", "jpg", "jpeg", "bmp", "tiff", "tif", "webp"}
        if not file or (str(file.get("format", "")).lower() != "pdf" and not image_format) or page_number < 1:
            errors.append({"file_id": selected.get("file_id"), "page": page_number, "reason": "page_not_renderable"})
            continue
        try:
            if _sha(Path(file["path"]).read_bytes()) != file["content_hash"]:
                raise SourceRejected("frozen_source_checksum_mismatch")
            if image_format:
                from PIL import Image
                import io
                with Image.open(file["path"]) as image:
                    if page_number > getattr(image, "n_frames", 1):
                        raise SourceRejected("image_frame_out_of_range")
                    image.seek(page_number - 1)
                    if image.width * image.height > 40_000_000:
                        raise SourceRejected("page_pixel_budget_exceeded")
                    buffer = io.BytesIO()
                    image.convert("RGB").save(buffer, format="PNG")
                    images.append({"file_id": file["file_id"], "page": page_number, "image_bytes": buffer.getvalue(), "content_hash": file["content_hash"]})
                continue
            document = pdfium.PdfDocument(file["path"])
            if page_number > len(document):
                raise ValueError("page_out_of_range")
            page = document[page_number - 1]
            try:
                width, height = page.get_size()
                if width * height * (dpi / 72) ** 2 > 40_000_000:
                    raise SourceRejected("page_pixel_budget_exceeded")
                bitmap = page.render(scale=dpi / 72)
                try:
                    pil = bitmap.to_pil()
                    import io
                    buffer = io.BytesIO()
                    pil.save(buffer, format="PNG")
                    images.append({"file_id": file["file_id"], "page": page_number, "image_bytes": buffer.getvalue(), "content_hash": file["content_hash"]})
                finally:
                    bitmap.close()
            finally:
                page.close()
                document.close()
        except Exception as exc:
            errors.append({"file_id": selected.get("file_id"), "page": page_number, "reason": type(exc).__name__})
    return {"images": images, "errors": errors}


def provider_evidence(snapshot: dict[str, Any], candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Register OCR transcription as a model source, never native truth."""
    files = {item["file_id"]: item for item in snapshot.get("files", [])}
    owner = {key: snapshot[key] for key in ("run_id", "mine_id", "environment", "rule_revision") if key in snapshot}
    evidence = []
    for candidate in candidates:
        file = files.get(candidate.get("file_id"))
        if not file:
            raise SourceRejected("provider candidate outside source snapshot")
        text = candidate.get("text", "")
        if not isinstance(text, str) or not text.strip():
            continue
        locator = {"page": candidate.get("page"), "provider": candidate.get("provider", "external"), "kind": "model_transcription"}
        identity = {**locator, "text_hash": _sha(text.encode("utf-8"))}
        evidence.append({**owner, "evidence_id": _stable_evidence(file["content_hash"], identity),
                         "file_id": file["file_id"], "document_id": file["file_id"], "document_version": file["content_hash"],
                         "content_hash": file["content_hash"], "locator": locator, "text": text, "quote": text,
                         "source_role": "statement", "origin_source_role": file.get("source_role", "statement"),
                         "unit_id": f"page:{candidate.get('page')}", "quality": "model_candidate",
                         "metadata": {"recognition_quality": "candidate", "quality_policy": "model_transcription_is_not_source_truth"}})
    return evidence


# Native parser migrated from compliance-service/parsers.py; no runtime dependency.
LEGACY_PARSER_SHA256 = "c8649dedfc1ea67f9c42d886318672e0df010156994e71ec4c3d33f756689770"
"""Local, read-only document extraction with explicit coverage and stable provenance.

No macros, hyperlinks, formulas or embedded objects are executed. OCR is deliberately
not inferred from an empty text layer: unsupported content remains visible as a gap.
"""

from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from time import perf_counter



def _block(result, text, kind, locator, metadata=None):
    identity = json.dumps(locator, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    result["blocks"].append({
        "block_id": sha256((result["content_hash"] + ":" + identity).encode("utf-8")).hexdigest(),
        "text": text, "kind": kind, "locator": locator, "metadata": metadata or {},
    })


def _coverage(result, unit_id, state, reason=None):
    item = {"unit_id": unit_id, "state": state}
    if reason:
        item["reason"] = reason
        issue = f"{unit_id}: {reason}"
        if issue not in result["issues"]:
            result["issues"].append(issue)
    result["coverage"].append(item)


def parse_document(path, *, file_id, progress_callback=None):
    """Extract a single immutable byte snapshot; file errors do not abort a batch."""
    path = Path(path)
    extension = path.suffix.lower().lstrip(".")
    result = {"file_id": file_id, "format": extension or "unknown", "content_hash": None,
              "blocks": [], "coverage": [], "issues": []}
    try:
        data = path.read_bytes()
        result["content_hash"] = sha256(data).hexdigest()
        if extension == "pdf":
            _pdf(data, result, progress_callback=progress_callback)
        elif extension == "docx":
            _docx(data, result)
        elif extension == "xlsx":
            _xlsx(data, result)
        elif extension in {"png", "jpg", "jpeg", "tiff", "tif", "bmp", "webp"}:
            _coverage(result, "document", "unreadable", "image_requires_ocr")
        else:
            _coverage(result, "document", "unsupported", f"unsupported_format:{extension or 'unknown'}")
    except Exception as exc:
        # Keep exception classes for diagnosis, without leaking local paths or binary content.
        _coverage(result, "document", "failed", f"parse_failed:{type(exc).__name__}")
    return result


def _emit_progress(callback, **event):
    if callback is None:
        return
    callback(event)


def _pdf(data, result, *, progress_callback=None):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data), strict=True)
    if reader.is_encrypted:
        _coverage(result, "document", "failed", "encrypted_pdf")
        return
    if not reader.pages:
        _coverage(result, "document", "empty", "pdf_has_no_pages")
    for number, page in enumerate(reader.pages, 1):
        started = perf_counter()
        unit = f"page:{number}"
        _emit_progress(progress_callback, stage="parse", event="page_started", file_id=result["file_id"],
                       page=number, page_total=len(reader.pages), status="running", issues=[])
        try:
            text = page.extract_text() or ""
            contents = page.get_contents()
            painted = _pdf_has_content(contents)
            has_images = _pdf_paints_images(contents, page.get("/Resources", {}), reader)
            if text.strip():
                _block(result, text, "page", {"page": number}, {"has_images": has_images})
            if has_images and text.strip():
                # Native text is still valid evidence. Image regions remain an
                # explicit gap and are eligible for OCR/VLM separately.
                _coverage(result, unit, "partial", "page_contains_image_requires_ocr")
            elif has_images:
                _coverage(result, unit, "unreadable", "page_image_requires_ocr")
            elif text.strip():
                _coverage(result, unit, "readable")
            elif painted:
                _coverage(result, unit, "unreadable", "nonempty_page_without_extractable_text")
            else:
                _coverage(result, unit, "empty")
        except Exception as exc:
            _coverage(result, unit, "failed", f"page_parse_failed:{type(exc).__name__}")
        _emit_progress(progress_callback, stage="parse", event="page_finished", file_id=result["file_id"],
                       page=number, page_total=len(reader.pages), status=result["coverage"][-1]["state"],
                       elapsed_ms=round((perf_counter() - started) * 1000),
                       issues=[result["coverage"][-1].get("reason")] if result["coverage"][-1].get("reason") else [])


def _pdf_paints_images(contents, resources, reader, seen=None):
    """Inspect only painted XObjects, including nested form streams.

    PDF generators commonly share one resource dictionary across every page;
    an unused image resource is not evidence that the page contains an image.
    """
    if contents is None:
        return False
    seen = set() if seen is None else seen
    resources = resources.get_object() if hasattr(resources, "get_object") else resources
    objects = resources.get("/XObject", {})
    objects = objects.get_object() if hasattr(objects, "get_object") else objects
    from pypdf.generic import ContentStream
    stream = contents if isinstance(contents, ContentStream) else ContentStream(contents, reader)
    for operands, operator in stream.operations:
        if operator == b"INLINE IMAGE":
            return True
        if operator != b"Do" or not operands:
            continue
        reference = objects.get(operands[0])
        if reference is None:
            continue
        key = (getattr(reference, "idnum", None), getattr(reference, "generation", None)) if hasattr(reference, "idnum") else id(reference)
        if key in seen:
            continue
        seen.add(key)
        obj = reference.get_object()
        if obj.get("/Subtype") == "/Image":
            return True
        if obj.get("/Subtype") == "/Form" and _pdf_paints_images(obj, obj.get("/Resources", resources), reader, seen):
            return True
    return False


def _pdf_has_content(contents):
    if contents is None:
        return False
    try:
        if hasattr(contents, "get_data"):
            return bool(contents.get_data().strip())
        return any(_pdf_has_content(item.get_object()) for item in contents)
    except Exception:
        return True


def _pdf_has_inline_image(contents):
    if contents is None:
        return False
    if hasattr(type(contents), "operations"):
        return any(operator == b"INLINE IMAGE" for _, operator in contents.operations)
    return any(_pdf_has_inline_image(item.get_object()) for item in contents)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _word_text(element):
    parts = []
    for node in element.iter():
        if node.tag in {_W + "t", _W + "delText"}:
            parts.append(node.text or "")
        elif node.tag == _W + "tab":
            parts.append("\t")
        elif node.tag in {_W + "br", _W + "cr"}:
            parts.append("\n")
    return "".join(parts)


def _word_gaps(element):
    tags = {node.tag.rsplit("}", 1)[-1] for node in element.iter()}
    reasons = []
    if tags & {"ins", "del", "moveFrom", "moveTo", "pPrChange", "rPrChange", "tblPrChange", "tcPrChange"}:
        reasons.append("tracked_changes_unresolved")
    if tags & {"drawing", "pict", "object"}:
        reasons.append("embedded_object_requires_extraction")
    if tags & {"fldChar", "fldSimple", "altChunk"}:
        reasons.append("dynamic_or_external_content_unresolved")
    return reasons


def _docx(data, result):
    from docx import Document

    document = Document(BytesIO(data))
    paragraph_number = table_number = 0
    for element in document.element.body:
        if element.tag == _W + "p":
            paragraph_number += 1
            unit = f"paragraph:{paragraph_number}"
            text = _word_text(element)
            gaps = _word_gaps(element)
            if text:
                _block(result, text, "paragraph", {"paragraph": paragraph_number}, {"gaps": gaps})
            _coverage(result, unit, "unreadable" if gaps else "readable" if text.strip() else "empty",
                      ";".join(gaps) or None)
        elif element.tag == _W + "tbl":
            table_number += 1
            rows = []
            for row in element.findall(_W + "tr"):
                properties = row.find(_W + "trPr")
                before_node = properties.find(_W + "gridBefore") if properties is not None else None
                before = int(before_node.get(_W + "val", "0")) if before_node is not None else 0
                cells, merges = [""] * before, []
                for cell in row.findall(_W + "tc"):
                    text = "\n".join(_word_text(p) for p in cell.findall(_W + "p"))
                    properties = cell.find(_W + "tcPr")
                    span_node = properties.find(_W + "gridSpan") if properties is not None else None
                    span = int(span_node.get(_W + "val", "1")) if span_node is not None else 1
                    merge_node = properties.find(_W + "vMerge") if properties is not None else None
                    merges.append({"column": len(cells) + 1, "span": span,
                                   "vertical": merge_node.get(_W + "val", "continue") if merge_node is not None else None})
                    cells.append(text)
                    cells.extend([""] * (span - 1))
                rows.append((row, cells, merges))
            headers = rows[0][1] if rows else []
            for row_number, (row, cells, merges) in enumerate(rows, 1):
                gaps = _word_gaps(row)
                if row.findall(".//" + _W + "tbl"):
                    gaps.append("nested_table_unresolved")
                _block(result, "\t".join(cells), "table_row", {"table": table_number, "row": row_number},
                       {"cells": cells, "headers": headers, "merges": merges, "gaps": gaps})
                _coverage(result, f"table:{table_number}:row:{row_number}",
                          "unreadable" if gaps else "readable" if any(c.strip() for c in cells) else "empty",
                          ";".join(gaps) or None)
        elif element.tag != _W + "sectPr":
            _coverage(result, f"body:{len(result['coverage']) + 1}", "unreadable", "unsupported_body_structure")
    # Headers, footers, notes and comments are not silently treated as inspected body text.
    for part in document.part.package.parts:
        name = str(part.partname)
        if any(token in name for token in ("/header", "/footer", "/footnotes", "/endnotes")):
            from lxml import etree
            root = etree.fromstring(part.blob)
            if _word_text(root).strip() or _word_gaps(root):
                _coverage(result, name, "unreadable", "non_body_content_unresolved")
    if not result["coverage"]:
        _coverage(result, "document", "empty")


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _json_value(value):
    return value.isoformat() if isinstance(value, (datetime, date, time)) else value


def _xlsx(data, result):
    from openpyxl import load_workbook

    formulas = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(data), data_only=True, keep_links=False)
    try:
        for sheet in formulas:
            cached_sheet = cached[sheet.title]
            merged = [str(area) for area in sheet.merged_cells.ranges]
            rows = []
            sheet_gaps = []
            if sheet._images or sheet._charts:
                sheet_gaps.append("sheet_drawing_requires_extraction")
            for row in sheet.iter_rows():
                # Preserve real cells, including numeric zero and formulas; avoid manufactured blank rows.
                if not any(cell.value is not None for cell in row):
                    continue
                cells, details, row_gaps = [], [], []
                for cell in row:
                    value = cell.value
                    formula = value if cell.data_type == "f" else None
                    cache = cached_sheet[cell.coordinate].value if formula is not None else None
                    if formula is not None:
                        value = cache
                        if cache is None:
                            row_gaps.append(f"formula_without_cached_value:{cell.coordinate}")
                        elif cached_sheet[cell.coordinate].data_type == "e":
                            row_gaps.append(f"formula_cached_error:{cell.coordinate}")
                    elif cell.data_type == "e":
                        row_gaps.append(f"cell_error:{cell.coordinate}")
                    cells.append(_cell_text(value))
                    merge_range = next((area for area in sheet.merged_cells.ranges if cell.coordinate in area), None)
                    details.append({
                        "coordinate": cell.coordinate, "value": _json_value(value),
                        "raw_value": _json_value(cell.value), "formula": formula,
                        "cached_value": _json_value(cache), "data_type": cell.data_type,
                        "number_format": cell.number_format,
                        "row_hidden": bool(sheet.row_dimensions[cell.row].hidden),
                        "column_hidden": _column_hidden(sheet, cell.column),
                        "merged_range": str(merge_range) if merge_range is not None else None,
                        "merged_anchor": sheet.cell(merge_range.min_row, merge_range.min_col).coordinate if merge_range is not None else None,
                    })
                rows.append((row[0].row, cells, details, row_gaps))
            headers = rows[0][1] if rows else []
            for number, cells, details, gaps in rows:
                locator = {"sheet": sheet.title, "row": number}
                _block(result, "\t".join(cells), "sheet_row", locator,
                       {"cells": cells, "headers": headers, "cell_details": details,
                        "sheet_state": sheet.sheet_state, "row_hidden": bool(sheet.row_dimensions[number].hidden),
                        "merged_ranges": merged, "gaps": gaps})
                if gaps:
                    sheet_gaps.extend(gaps)
                    _coverage(result, f"sheet:{sheet.title}:row:{number}", "unreadable", ";".join(gaps))
            _coverage(result, f"sheet:{sheet.title}", "unreadable" if sheet_gaps else "readable" if rows else "empty",
                      ";".join(sheet_gaps) or None)
    finally:
        formulas.close()
        cached.close()


def _column_hidden(sheet, number):
    # Excel can store one grouped dimension for several columns.
    return any(dimension.hidden and (dimension.min or 0) <= number <= (dimension.max or 0)
               for dimension in sheet.column_dimensions.values())
